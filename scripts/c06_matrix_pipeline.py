#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PARSER_NAME = "c06_matrix_v3"
PARSER_VERSION = "1.0.0"
EXPECTED_SHEET = "Matriz de Contingência"
EXPECTED_HEADERS = {
    2: "#",
    3: "Cenário",
    4: "Gatilho",
    5: "Acionamento",
    6: "Protocolo",
    7: "Área responsável",
    8: "Dono",
    9: "Áreas impactadas",
    10: "SLA-alvo",
    11: "Ferramenta",
    12: "Acompanhamento e visibilidade",
    13: "Participantes reunião de validação",
    14: "Mapeamento",
}
CANONICAL_FIELDS = [
    "number",
    "code",
    "name",
    "trigger",
    "activation",
    "protocol",
    "responsible_area",
    "owner",
    "impacted_areas",
    "sla_target",
    "tool",
    "monitoring_visibility",
    "validation_participants",
    "mapping",
    "source_row",
    "source_file",
    "source_sheet",
    "source_sha256",
]
EXPECTED_OWNER_MAP = {
    1: "Daniel Garcia",
    2: "Daniel Garcia",
    3: "Daniel Garcia",
    4: "Jiane Rodrigues",
    5: "Jiane Rodrigues",
    6: "Jiane Rodrigues",
    7: "Daniel Garcia",
    8: "Jiane Rodrigues",
    9: "Renato de Paulo",
    10: "Daniel Garcia",
    11: "Daniel Garcia",
}
MANDATORY_TEXT_FIELDS = [
    "name",
    "trigger",
    "activation",
    "protocol",
    "responsible_area",
    "owner",
    "impacted_areas",
    "sla_target",
    "tool",
    "monitoring_visibility",
    "validation_participants",
    "mapping",
]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def sha256_json(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def parse_xlsx(path: Path, expected_sha: str | None = None) -> dict[str, Any]:
    source_sha = sha256_file(path)
    if expected_sha and source_sha != expected_sha:
        raise ValueError(f"source SHA mismatch: expected {expected_sha}, got {source_sha}")

    wb = load_workbook(path, data_only=False, read_only=False)
    if EXPECTED_SHEET not in wb.sheetnames:
        raise ValueError(f"missing sheet: {EXPECTED_SHEET}")
    ws = wb[EXPECTED_SHEET]

    for col, expected in EXPECTED_HEADERS.items():
        actual = text(ws.cell(4, col).value)
        if actual != expected:
            raise ValueError(f"header mismatch at row 4 col {col}: expected {expected!r}, got {actual!r}")

    starts: list[int] = []
    for row in range(5, ws.max_row + 1):
        value = ws.cell(row, 2).value
        if isinstance(value, int):
            starts.append(row)
        elif isinstance(value, float) and value.is_integer():
            starts.append(row)

    records: list[dict[str, Any]] = []
    for i, start in enumerate(starts):
        number = int(ws.cell(start, 2).value)
        end = starts[i + 1] - 1 if i + 1 < len(starts) else ws.max_row
        protocol_lines = [text(ws.cell(r, 6).value) for r in range(start, end + 1) if text(ws.cell(r, 6).value)]
        record = {
            "number": number,
            "code": f"SAFRA-{number:02d}",
            "name": text(ws.cell(start, 3).value),
            "trigger": text(ws.cell(start, 4).value),
            "activation": text(ws.cell(start, 5).value),
            "protocol": "\n".join(protocol_lines),
            "responsible_area": text(ws.cell(start, 7).value),
            "owner": text(ws.cell(start, 8).value),
            "impacted_areas": text(ws.cell(start, 9).value),
            "sla_target": text(ws.cell(start, 10).value),
            "tool": text(ws.cell(start, 11).value),
            "monitoring_visibility": text(ws.cell(start, 12).value),
            "validation_participants": text(ws.cell(start, 13).value),
            "mapping": text(ws.cell(start, 14).value),
            "source_row": start,
            "source_file": "EDB06 - Matriz Contingencia v3.xlsx",
            "source_sheet": EXPECTED_SHEET,
            "source_sha256": source_sha,
        }
        records.append(record)

    payload = {
        "pipeline_stage": "staging",
        "parser": {"name": PARSER_NAME, "version": PARSER_VERSION},
        "source": {
            "logical_name": "EDB06 - Matriz Contingencia v3.xlsx",
            "materialized_name": path.name,
            "sheet": EXPECTED_SHEET,
            "sha256": source_sha,
            "workbook_sheets": wb.sheetnames,
        },
        "record_count": len(records),
        "records": records,
    }
    payload["staging_sha256"] = sha256_json(records)
    return payload


def validate(staging: dict[str, Any]) -> dict[str, Any]:
    records = staging["records"]
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []

    numbers = [r["number"] for r in records]
    codes = [r["code"] for r in records]
    if len(records) != 11:
        errors.append({"check": "record_count", "expected": 11, "observed": len(records)})
    if numbers != list(range(1, 12)):
        errors.append({"check": "numbers", "expected": list(range(1, 12)), "observed": numbers})
    if len(set(codes)) != len(codes):
        errors.append({"check": "unique_codes", "observed": codes})

    for r in records:
        for field in MANDATORY_TEXT_FIELDS:
            if not r[field].strip():
                errors.append({"check": "mandatory_field", "scenario": r["code"], "field": field})
        expected_owner = EXPECTED_OWNER_MAP.get(r["number"])
        if r["owner"] != expected_owner:
            errors.append({
                "check": "owner_map",
                "scenario": r["code"],
                "expected": expected_owner,
                "observed": r["owner"],
            })
        if r["mapping"] not in {"EDB05", "EDB06"}:
            errors.append({"check": "mapping", "scenario": r["code"], "observed": r["mapping"]})
        protocol_steps = [line for line in r["protocol"].split("\n") if re.match(r"^\d+\)", line.strip())]
        if len(protocol_steps) != 5:
            errors.append({"check": "protocol_steps", "scenario": r["code"], "expected": 5, "observed": len(protocol_steps)})

    # Explicitly open business decisions; these are not parser errors.
    open_issue_patterns = [
        ("GI-SAFRA-002", "SAFRA-02", r"\bX h\b", "threshold in hours is unresolved"),
        ("GI-SAFRA-002", "SAFRA-04", r"\bX min\b", "threshold in minutes is unresolved"),
        ("GI-SAFRA-002", "SAFRA-10", r"acima do limite", "lead-time/queue limit is unresolved"),
        ("GI-SAFRA-002", "SAFRA-11", r"limiar planejado", "capacity threshold is unresolved"),
        ("GI-SAFRA-003", "SAFRA-09", r"saldo.*mínimo|abaixo do mínimo", "curve-A minimum source/rule is unresolved"),
    ]
    by_code = {r["code"]: r for r in records}
    for issue, code, pattern, description in open_issue_patterns:
        haystack = "\n".join(str(by_code[code].get(k, "")) for k in ("trigger", "activation", "protocol", "sla_target"))
        if re.search(pattern, haystack, flags=re.IGNORECASE):
            warnings.append({"governance_issue": issue, "scenario": code, "description": description})

    result = {
        "pipeline_stage": "validation",
        "parser": staging["parser"],
        "source_sha256": staging["source"]["sha256"],
        "staging_sha256": staging["staging_sha256"],
        "record_count": len(records),
        "status": "PASS" if not errors else "FAIL",
        "errors": errors,
        "warnings": warnings,
        "criticality_inferred": False,
        "structured_sla_inferred": False,
    }
    result["validation_sha256"] = sha256_json(result)
    return result


def diff_staging(staging: dict[str, Any], baseline: dict[str, Any]) -> dict[str, Any]:
    left = {r["code"]: r for r in staging["records"]}
    baseline_records = baseline.get("records", baseline.get("canonical_records", baseline))
    if not isinstance(baseline_records, list):
        raise ValueError("baseline must contain a list under records/canonical_records or be a list")
    right = {r["code"]: r for r in baseline_records}

    added = sorted(set(left) - set(right))
    removed = sorted(set(right) - set(left))
    changed: list[dict[str, Any]] = []
    for code in sorted(set(left) & set(right)):
        fields = []
        for field in CANONICAL_FIELDS:
            if left[code].get(field) != right[code].get(field):
                fields.append({"field": field, "from": right[code].get(field), "to": left[code].get(field)})
        if fields:
            changed.append({"code": code, "fields": fields})

    result = {
        "pipeline_stage": "preview_diff",
        "parser": staging["parser"],
        "source_sha256": staging["source"]["sha256"],
        "staging_sha256": staging["staging_sha256"],
        "baseline_record_count": len(right),
        "staging_record_count": len(left),
        "added": added,
        "removed": removed,
        "changed": changed,
        "difference_count": len(added) + len(removed) + sum(len(x["fields"]) for x in changed),
        "status": "NO_DIFF" if not added and not removed and not changed else "DIFF_REQUIRES_HUMAN_REVIEW",
    }
    result["diff_sha256"] = sha256_json(result)
    return result


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=False) + "\n", encoding="utf-8")


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--xlsx", type=Path, required=True)
    p.add_argument("--expected-sha")
    p.add_argument("--staging", type=Path, required=True)
    p.add_argument("--validation", type=Path, required=True)
    p.add_argument("--baseline", type=Path)
    p.add_argument("--diff", type=Path)
    args = p.parse_args()

    staging = parse_xlsx(args.xlsx, args.expected_sha)
    validation = validate(staging)
    write_json(args.staging, staging)
    write_json(args.validation, validation)

    if args.baseline:
        baseline = json.loads(args.baseline.read_text(encoding="utf-8"))
        diff = diff_staging(staging, baseline)
        if not args.diff:
            raise ValueError("--diff is required when --baseline is supplied")
        write_json(args.diff, diff)
        if diff["status"] != "NO_DIFF":
            return 3

    return 0 if validation["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())